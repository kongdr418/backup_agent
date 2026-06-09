"""Main PPT generation pipeline orchestrator.

Runs 5 stages sequentially:
1. Content Planning (topic → manuscript)
2. Design Strategy (manuscript → design_spec)
3. SVG Generation (per-page SVG with critic)
4. SVG Finalization (post-processing)
5. PPTX Export (SVG → DrawingML → PPTX)
"""

from __future__ import annotations

import asyncio
import json
import time
from collections.abc import AsyncIterator
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from ppt_engine.agents.content_planner import _split_manuscript_pages, plan_content
from ppt_engine.agents.design_strategist import create_design_spec
from ppt_engine.agents.svg_executor import generate_svg_pages
from ppt_engine.config import DESIGN_STYLES, WORKSPACES_DIR, get_deepseek_api_key, VISUAL_CRITIC_ENABLED, VISUAL_CRITIC_MODEL, DEEP_RESEARCH_ENABLED, DEEP_RESEARCH_QUALITY_THRESHOLD, DEEP_RESEARCH_MAX_ATTEMPTS
from ppt_engine.llm.deepseek_provider import DeepSeekProvider
from ppt_engine.llm.anthropic_provider import AnthropicProvider
from ppt_engine.llm.base import LLMProvider
from ppt_engine.logger import get_logger


logger = get_logger(__name__)


@dataclass
class PipelineEvent:
    """Event yielded by the pipeline for real-time progress tracking."""
    stage: str        # "content_planning", "design", "svg_generation", "finalize", "export"
    status: str       # "started", "progress", "complete", "error"
    message: str
    progress: float   # 0.0 ~ 1.0
    data: dict | None = field(default_factory=dict)

    def to_dict(self) -> dict:
        d = {
            "stage": self.stage,
            "status": self.status,
            "message": self.message,
            "progress": round(self.progress, 3),
        }
        if self.data:
            d.update(self.data)
        return d


class PPTPipeline:
    """SVG-based PPT generation pipeline."""

    @staticmethod
    def _create_llm(api_key: str, base_url: str | None, provider_name: str) -> LLMProvider:
        """根据 base_url 选择 LLM provider。"""
        is_anthropic = (
            "/api/anthropic" in (base_url or "")
            or provider_name == "zhipu"
        )
        if is_anthropic:
            return AnthropicProvider(api_key=api_key, base_url=base_url, provider_name=provider_name)
        return DeepSeekProvider(api_key=api_key, base_url=base_url, provider_name=provider_name)

    async def generate(
        self,
        topic: str,
        *,
        provider: str = "deepseek",
        model: str = "deepseek-v4-flash",
        api_key: str | None = None,
        base_url: str | None = None,
        user_id: str = "anonymous",
        language: str = "zh",
        num_slides: int | None = None,
        detail_level: str = "normal",
        style: str = "education",
        canvas_format: str = "ppt169",
        style_overrides: dict | None = None,
        deep_research: bool = False,
        visual_critic: bool = False,
        repair_enabled: bool = False,
        template_id: str | None = None,
        notes: str | None = None,
    ) -> AsyncIterator[PipelineEvent]:
        """Generate a PPT from a course topic.

        Yields PipelineEvent objects for real-time progress tracking.
        """
        # Resolve API key
        if not api_key:
            api_key = get_deepseek_api_key()
        if not api_key:
            yield PipelineEvent("init", "error", "未配置 API Key", 0.0)
            return

        # Create LLM provider based on base_url
        llm: LLMProvider = self._create_llm(api_key, base_url, provider)

        # Create user-scoped job workspace
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        job_id = f"course_ppt_{timestamp}"
        workspace = WORKSPACES_DIR
        if user_id != 'anonymous':
            workspace = WORKSPACES_DIR / 'users' / user_id
        workspace.mkdir(parents=True, exist_ok=True)
        project_dir = workspace / job_id
        project_dir.mkdir(parents=True, exist_ok=True)
        pipeline_start = time.monotonic()

        style_info = DESIGN_STYLES.get(style, DESIGN_STYLES["education"])

        # Load template if specified
        tmpl = None
        if template_id:
            from ppt_engine.template_manager import load_template
            tmpl = load_template(template_id)
            if tmpl is None:
                yield PipelineEvent("init", "error", f"模板 '{template_id}' 不存在", 0.0)
                return

        yield PipelineEvent(
            "init", "started",
            f"开始生成: {topic}",
            0.0,
            {"job_id": job_id, "project_dir": str(project_dir)},
        )
        logger.info(
            "[PPT-PERF] job=%s stage=init status=started topic=%s slides=%s model=%s template=%s repair_enabled=%s",
            job_id,
            topic,
            num_slides,
            model,
            template_id or "",
            repair_enabled,
        )

        # Save metadata
        meta = {
            "job_id": job_id,
            "topic": topic,
            "language": language,
            "style": style,
            "model": model,
            "requested_num_slides": num_slides,
            "created_at": timestamp,
        }
        (project_dir / "metadata.json").write_text(
            json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
        )

        try:
            # ── Stage 1: Content Planning (0% → 15%) ──
            stage_start = time.monotonic()
            logger.info("[PPT-PERF] job=%s stage=content_planning status=started", job_id)
            yield PipelineEvent("content_planning", "started", "正在规划课程内容...", 0.05)

            if deep_research:
                from ppt_engine.agents.research_agent import run_deep_research
                manuscript = None
                async for event in run_deep_research(
                    topic, llm, model,
                    language=language,
                    num_slides=num_slides,
                    detail_level=detail_level,
                    debug_dir=project_dir / "debug",
                ):
                    if event.stage == "research_complete" and event.data:
                        manuscript = event.data.get("manuscript")
                    else:
                        yield PipelineEvent(
                            event.stage, event.status, event.message, event.progress, event.data
                        )
                if manuscript is None:
                    yield PipelineEvent("error", "error", "深度研究未生成手稿", 0.0)
                    return
            else:
                manuscript = await plan_content(
                    topic, llm, model,
                    instruction=(notes or ""),  # 备注透传；None/空 → ""
                    language=language,
                    num_slides=num_slides,
                    detail_level=detail_level,
                )
            (project_dir / "manuscript.md").write_text(manuscript, encoding="utf-8")

            slide_count = len(_split_manuscript_pages(manuscript))
            yield PipelineEvent(
                "content_planning", "complete",
                f"课程规划完成: {slide_count} 页幻灯片",
                0.15,
                {"slide_count": slide_count},
            )
            logger.info(
                "[PPT-PERF] job=%s stage=content_planning status=complete elapsed=%.2fs slides=%s",
                job_id,
                time.monotonic() - stage_start,
                slide_count,
            )

            # ── Stage 2: Design Strategy (15% → 30%) ──
            stage_start = time.monotonic()
            logger.info("[PPT-PERF] job=%s stage=design status=started", job_id)
            yield PipelineEvent("design", "started", "正在生成设计规范...", 0.15)

            if tmpl and len(tmpl.design_spec) >= 500:
                design_spec = tmpl.design_spec
                yield PipelineEvent("design", "complete", "使用模板设计规范", 0.25)
            else:
                design_spec = await create_design_spec(
                    manuscript, llm, model,
                    canvas_format=canvas_format,
                    style=style,
                    language=language,
                    detail_level=detail_level,
                    style_overrides=style_overrides,
                )
                yield PipelineEvent("design", "complete", "设计规范生成完成", 0.30)
            (project_dir / "design_spec.md").write_text(design_spec, encoding="utf-8")
            logger.info(
                "[PPT-PERF] job=%s stage=design status=complete elapsed=%.2fs source=%s",
                job_id,
                time.monotonic() - stage_start,
                "template" if tmpl and len(tmpl.design_spec) >= 500 else "llm",
            )

            # ── Stage 3: SVG Generation (30% → 75%) ──
            stage_start = time.monotonic()
            logger.info("[PPT-PERF] job=%s stage=svg_generation status=started slides=%s", job_id, slide_count)
            yield PipelineEvent("svg_generation", "started", "正在逐页生成 SVG...", 0.30)

            # Load template context for SVG generation
            template_context = None
            template_skeletons = None
            if tmpl:
                from ppt_engine.template_manager import build_template_context, build_template_skeletons
                template_context = build_template_context(tmpl)
                template_skeletons = build_template_skeletons(tmpl)

            svg_pages: list[tuple[int, str]] = []
            page_progress_base = 0.30
            page_progress_range = 0.45
            completed_pages = 0

            async for page_num, svg_content in generate_svg_pages(
                design_spec, manuscript, project_dir, llm, model,
                style=style,
                language=language,
                detail_level=detail_level,
                template_context=template_context,
                template_svgs=template_skeletons,
                repair_enabled=repair_enabled,
            ):
                svg_pages.append((page_num, svg_content))
                completed_pages += 1
                progress = page_progress_base + (completed_pages / slide_count) * page_progress_range
                yield PipelineEvent(
                    "svg_generation", "progress",
                    f"第{page_num}页生成完成",
                    progress,
                    {
                        "page": page_num,
                        "total_slides": slide_count,
                        "svg": svg_content,
                    },
                )

            yield PipelineEvent(
                "svg_generation", "complete",
                f"全部 {len(svg_pages)} 页 SVG 生成完成",
                0.75,
            )
            logger.info(
                "[PPT-PERF] job=%s stage=svg_generation status=complete elapsed=%.2fs slides=%s",
                job_id,
                time.monotonic() - stage_start,
                len(svg_pages),
            )

            # ── Stage 4: SVG Finalization (75% → 85%) ──
            stage_start = time.monotonic()
            logger.info("[PPT-PERF] job=%s stage=finalize status=started", job_id)
            yield PipelineEvent("finalize", "started", "正在后处理 SVG...", 0.75)

            svg_output_dir = project_dir / "svg_output"
            svg_final_dir = project_dir / "svg_final"
            svg_final_dir.mkdir(parents=True, exist_ok=True)

            # Run finalize pipeline on each SVG
            from ppt_engine.finalize.finalize import finalize_svg_dir
            try:
                finalize_svg_dir(svg_output_dir, svg_final_dir)
            except Exception as e:
                # If finalize fails, copy raw SVGs as fallback
                import shutil
                for svg_file in svg_output_dir.glob("*.svg"):
                    shutil.copy2(svg_file, svg_final_dir / svg_file.name)

            yield PipelineEvent("finalize", "complete", "SVG 后处理完成", 0.85)
            logger.info(
                "[PPT-PERF] job=%s stage=finalize status=complete elapsed=%.2fs",
                job_id,
                time.monotonic() - stage_start,
            )

            # ── Stage 5: PPTX Export (85% → 100%) ──
            stage_start = time.monotonic()
            logger.info("[PPT-PERF] job=%s stage=export status=started", job_id)
            yield PipelineEvent("export", "started", "正在导出 PPTX...", 0.85)

            exports_dir = project_dir / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)
            pptx_filename = f"presentation_{timestamp}.pptx"
            pptx_path = exports_dir / pptx_filename

            try:
                from ppt_engine.svg_to_pptx.builder import create_pptx
                svg_files = sorted(svg_final_dir.glob("*.svg"))
                # Parse manuscript into per-slide notes
                notes = _parse_manuscript_to_notes(manuscript, svg_files)
                create_pptx(
                    svg_files=svg_files,
                    output_path=pptx_path,
                    canvas_format=canvas_format,
                    notes=notes,
                )
            except Exception as e:
                yield PipelineEvent(
                    "export", "progress",
                    f"PPTX 原生导出失败 ({e})，使用 PNG 降级方案...",
                    0.90,
                )
                # Fallback: rasterize SVGs to PNG and embed
                self._fallback_png_export(svg_final_dir, pptx_path, canvas_format)

            yield PipelineEvent(
                "export", "complete",
                "PPT 生成完成!",
                1.0,
                {
                    "output_path": str(pptx_path),
                    "pptx_filename": pptx_filename,
                    "job_id": job_id,
                    "total_slides": len(svg_pages),
                },
            )
            logger.info(
                "[PPT-PERF] job=%s stage=export status=complete elapsed=%.2fs output=%s",
                job_id,
                time.monotonic() - stage_start,
                pptx_path,
            )
            logger.info(
                "[PPT-PERF] job=%s stage=total status=complete elapsed=%.2fs",
                job_id,
                time.monotonic() - pipeline_start,
            )

        except Exception as e:
            logger.exception(
                "[PPT-PERF] job=%s stage=total status=error elapsed=%.2fs",
                job_id,
                time.monotonic() - pipeline_start,
            )
            yield PipelineEvent("error", "error", f"生成失败: {e}", 0.0, {"error": str(e)})

    def _fallback_png_export(self, svg_dir: Path, pptx_path: Path, canvas_format: str) -> None:
        """Fallback: rasterize SVGs to PNG and embed in PPTX."""
        from pptx import Presentation
        from pptx.util import Inches
        from svglib.svglib import svg2rlg
        from reportlab.graphics import renderPM
        from PIL import Image
        import io

        from ppt_engine.config import CANVAS_FORMATS

        fmt = CANVAS_FORMATS.get(canvas_format, CANVAS_FORMATS["ppt169"])
        prs = Presentation()
        prs.slide_width = Inches(fmt["width"] / 96)
        prs.slide_height = Inches(fmt["height"] / 96)

        svg_files = sorted(svg_dir.glob("*.svg"))
        for svg_file in svg_files:
            try:
                drawing = svg2rlg(str(svg_file))
                if drawing is None:
                    continue
                png_data = renderPM.drawToString(drawing, fmt="PNG")
                img = Image.open(io.BytesIO(png_data))
                img_path = svg_dir / f"{svg_file.stem}.png"
                img.save(str(img_path), "PNG")

                slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
                slide.shapes.add_picture(
                    str(img_path),
                    Inches(0), Inches(0),
                    prs.slide_width, prs.slide_height,
                )
            except Exception:
                continue

        prs.save(str(pptx_path))


def _parse_manuscript_to_notes(manuscript: str, svg_files: list[Path]) -> dict[str, str]:
    """Split manuscript into per-slide notes (already plain text, just split by ---)."""
    import re
    notes = {}
    sections = re.split(r"\n---\n", manuscript.strip())
    for i, svg_file in enumerate(svg_files):
        if i < len(sections):
            text = sections[i].strip()
            if text:
                notes[svg_file.stem] = text
    return notes
